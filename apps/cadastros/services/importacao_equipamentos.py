from io import BytesIO

from openpyxl import load_workbook

from apps.cadastros.models import (
    ClasseOperacional,
    Empresa,
    Equipamento,
    StatusEquipamento,
)

COLUNAS_OBRIGATORIAS = [
    "Frota",
    "Empresa",
    "Classe Operacional",
    "Status Operacional",
    "Descrição",
]

COLUNAS_OPCIONAIS = [
    "Placa",
    "Marca",
    "Modelo",
    "Ano Fabricação",
    "Ano Modelo",
    "Renavam",
    "Chassi",
    "Cor",
    "Observação",
]


def texto(valor):
    if valor is None:
        return ""

    return str(valor).strip()


def normalizar_ano(valor, nome_coluna):
    if valor is None or texto(valor) == "":
        return None, None

    try:
        numero = float(valor)

        if not numero.is_integer():
            raise ValueError

        ano = int(numero)

    except (ValueError, TypeError):
        return None, f"{nome_coluna} deve ser um número inteiro."

    if ano < 1900 or ano > 2155:
        return None, f"{nome_coluna} está fora do intervalo permitido."

    return ano, None


def analisar_arquivo(arquivo):
    workbook = load_workbook(
        filename=BytesIO(arquivo.read()),
        read_only=True,
        data_only=True,
    )

    if "Equipamentos" not in workbook.sheetnames:
        raise ValueError("A planilha precisa conter uma aba chamada 'Equipamentos'.")

    ws = workbook["Equipamentos"]

    linhas = ws.iter_rows(values_only=True)

    cabecalho_original = next(linhas, None)

    if not cabecalho_original:
        raise ValueError("A planilha está vazia.")

    cabecalho = [texto(valor) for valor in cabecalho_original]

    if len(cabecalho) != len(set(cabecalho)):
        raise ValueError("Existem nomes de colunas duplicados na planilha.")

    faltantes = [coluna for coluna in COLUNAS_OBRIGATORIAS if coluna not in cabecalho]

    if faltantes:
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(faltantes))

    colunas_permitidas = COLUNAS_OBRIGATORIAS + COLUNAS_OPCIONAIS

    desconhecidas = [
        coluna for coluna in cabecalho if coluna and coluna not in colunas_permitidas
    ]

    if desconhecidas:
        raise ValueError("Colunas não reconhecidas: " + ", ".join(desconhecidas))

    indices = {coluna: cabecalho.index(coluna) for coluna in cabecalho if coluna}

    empresas = {}

    for empresa in Empresa.objects.all():
        nomes = {
            texto(empresa.razao_social).casefold(),
        }

        if empresa.nome_fantasia:
            nomes.add(texto(empresa.nome_fantasia).casefold())

        for nome in nomes:
            empresas[nome] = empresa

    classes = {
        texto(obj.nome).casefold(): obj for obj in ClasseOperacional.objects.all()
    }

    status_map = {
        texto(obj.nome).casefold(): obj for obj in StatusEquipamento.objects.all()
    }

    frotas_existentes = {
        texto(valor).casefold()
        for valor in Equipamento.objects.values_list(
            "frota",
            flat=True,
        )
    }

    frotas_arquivo = set()
    registros = []
    numero_linha = 1

    for linha in linhas:
        numero_linha += 1

        valores = list(linha)

        if not any(valor is not None and texto(valor) for valor in valores):
            continue

        def obter(nome):
            indice = indices.get(nome)

            if indice is None or indice >= len(valores):
                return ""

            return texto(valores[indice])

        frota = obter("Frota")
        nome_empresa = obter("Empresa")
        nome_classe = obter("Classe Operacional")
        nome_status = obter("Status Operacional")
        descricao = obter("Descrição")

        erros = []

        if not frota:
            erros.append("Frota não informada.")

        if not nome_empresa:
            erros.append("Empresa não informada.")

        if not nome_classe:
            erros.append("Classe operacional não informada.")

        if not nome_status:
            erros.append("Status operacional não informado.")

        if not descricao:
            erros.append("Descrição não informada.")

        empresa = empresas.get(nome_empresa.casefold())
        classe = classes.get(nome_classe.casefold())
        status = status_map.get(nome_status.casefold())

        if nome_empresa and not empresa:
            erros.append(f"Empresa '{nome_empresa}' não encontrada.")

        if nome_classe and not classe:
            erros.append(f"Classe operacional '{nome_classe}' não encontrada.")

        if nome_status and not status:
            erros.append(f"Status operacional '{nome_status}' não encontrado.")

        frota_normalizada = frota.casefold()

        if frota:
            if frota_normalizada in frotas_existentes:
                erros.append(f"A frota '{frota}' já está cadastrada.")

            if frota_normalizada in frotas_arquivo:
                erros.append(f"A frota '{frota}' está duplicada na planilha.")

            frotas_arquivo.add(frota_normalizada)

        ano_fabricacao, erro_fabricacao = normalizar_ano(
            obter("Ano Fabricação"),
            "Ano de fabricação",
        )

        ano_modelo, erro_modelo = normalizar_ano(
            obter("Ano Modelo"),
            "Ano modelo",
        )

        if erro_fabricacao:
            erros.append(erro_fabricacao)

        if erro_modelo:
            erros.append(erro_modelo)

        dados_registro = {
            "linha": numero_linha,
            "frota": frota,
            "empresa_id": empresa.pk if empresa else None,
            "empresa_nome": nome_empresa,
            "classe_operacional_id": classe.pk if classe else None,
            "classe_operacional_nome": nome_classe,
            "status_operacional_id": status.pk if status else None,
            "status_operacional_nome": nome_status,
            "descricao": descricao,
            "placa": obter("Placa"),
            "marca": obter("Marca"),
            "modelo": obter("Modelo"),
            "ano_fabricacao": ano_fabricacao,
            "ano_modelo": ano_modelo,
            "renavam": obter("Renavam"),
            "chassi": obter("Chassi"),
            "cor": obter("Cor"),
            "observacao": obter("Observação"),
            "erros": erros,
        }

        registros.append(dados_registro)

    workbook.close()

    validos = sum(1 for registro in registros if not registro["erros"])

    invalidos = len(registros) - validos

    return {
        "registros": registros,
        "resumo": {
            "total": len(registros),
            "validos": validos,
            "invalidos": invalidos,
        },
    }


def executar_importacao(dados):
    from django.db import transaction

    @transaction.atomic
    def importar():
        criados = 0

        for registro in dados["registros"]:
            if registro["erros"]:
                continue

            equipamento = Equipamento(
                frota=registro["frota"],
                empresa_id=registro["empresa_id"],
                classe_operacional_id=registro["classe_operacional_id"],
                status_operacional_id=registro["status_operacional_id"],
                descricao=registro["descricao"],
                placa=registro["placa"],
                marca=registro["marca"],
                modelo=registro["modelo"],
                ano_fabricacao=registro["ano_fabricacao"],
                ano_modelo=registro["ano_modelo"],
                renavam=registro["renavam"],
                chassi=registro["chassi"],
                cor=registro["cor"],
                observacao=registro["observacao"],
                ativo=True,
            )

            equipamento.save()
            criados += 1

        return criados

    return importar()
