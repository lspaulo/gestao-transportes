from collections import defaultdict
from datetime import date, datetime
from io import BytesIO

from django.db import transaction
from openpyxl import load_workbook

from apps.cadastros.models import (
    ContaBancariaFuncionario,
    Empresa,
    Funcao,
    Funcionario,
)
from apps.cadastros.models.bancos import Banco


class ImportacaoFuncionarioError(Exception):
    pass


def texto(valor):
    if valor is None:
        return ""

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    return str(valor).strip()


def normalizar_cpf(valor):
    """
    Mantém somente os números do CPF.
    """
    return "".join(ch for ch in texto(valor) if ch.isdigit())


def normalizar_data(valor):
    if valor is None or valor == "":
        return None

    if isinstance(valor, datetime):
        return valor.date().isoformat()

    if isinstance(valor, date):
        return valor.isoformat()

    valor = texto(valor)

    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return (
                datetime.strptime(
                    valor,
                    formato,
                )
                .date()
                .isoformat()
            )
        except ValueError:
            continue

    raise ValueError(f"Data inválida: {valor}. Use o formato DD/MM/AAAA.")


def normalizar(valor):
    return texto(valor).strip().casefold()


def valor_bool(valor):
    valor = normalizar(valor)

    return valor in {
        "sim",
        "s",
        "true",
        "1",
        "x",
        "yes",
    }


def obter_headers(worksheet):
    headers = {}

    for cell in worksheet[1]:
        if cell.value:
            headers[normalizar(cell.value)] = cell.column

    return headers


def valor_coluna(row, headers, nome):
    coluna = headers.get(normalizar(nome))

    if not coluna:
        return ""

    return row[coluna - 1].value


def validar_colunas(headers, obrigatorias):
    faltantes = []

    for coluna in obrigatorias:
        if normalizar(coluna) not in headers:
            faltantes.append(coluna)

    if faltantes:
        raise ImportacaoFuncionarioError(
            "Colunas obrigatórias ausentes: " + ", ".join(faltantes)
        )


def mapa_empresas():
    resultado = {}

    for empresa in Empresa.objects.all():
        resultado[normalizar(empresa.razao_social)] = empresa

        if empresa.nome_fantasia:
            resultado[normalizar(empresa.nome_fantasia)] = empresa

    return resultado


def mapa_funcoes():
    return {normalizar(funcao.nome): funcao for funcao in Funcao.objects.all()}


def mapa_bancos():
    resultado = {}

    for codigo, descricao in Banco.choices:
        resultado[normalizar(codigo)] = codigo
        resultado[normalizar(descricao)] = codigo
        resultado[normalizar(f"{codigo} - {descricao}")] = codigo

    return resultado


def mapa_tipo_conta():
    resultado = {}

    for codigo, descricao in ContaBancariaFuncionario.TipoConta.choices:
        resultado[normalizar(codigo)] = codigo
        resultado[normalizar(descricao)] = codigo

    return resultado


def mapa_tipo_pix():
    resultado = {}

    for codigo, descricao in ContaBancariaFuncionario.TipoChavePix.choices:
        resultado[normalizar(codigo)] = codigo
        resultado[normalizar(descricao)] = codigo

    return resultado


def ler_funcionarios(worksheet):
    headers = obter_headers(worksheet)

    validar_colunas(
        headers,
        [
            "Nome",
            "CPF",
            "Empresa",
        ],
    )

    empresas = mapa_empresas()
    funcoes = mapa_funcoes()

    registros = []
    cpfs_excel = {}

    for numero_linha, row in enumerate(
        worksheet.iter_rows(min_row=2),
        start=2,
    ):
        nome = texto(valor_coluna(row, headers, "Nome"))

        cpf_original = texto(valor_coluna(row, headers, "CPF"))

        cpf = normalizar_cpf(cpf_original)

        empresa_nome = texto(valor_coluna(row, headers, "Empresa"))

        funcao_nome = texto(valor_coluna(row, headers, "Função"))

        # Ignora linhas totalmente vazias
        if not nome and not cpf and not empresa_nome:
            continue

        erros = []

        if not nome:
            erros.append("Nome não informado.")

        if not cpf:
            erros.append("CPF não informado.")
        elif len(cpf) != 11:
            erros.append("CPF deve possuir 11 dígitos.")

        if cpf in cpfs_excel:
            erros.append(
                f"CPF duplicado no arquivo. Também aparece na linha {cpfs_excel[cpf]}."
            )
        else:
            if cpf:
                cpfs_excel[cpf] = numero_linha

        empresa = empresas.get(normalizar(empresa_nome))

        if not empresa:
            erros.append(f"Empresa '{empresa_nome}' não encontrada.")

        funcao = None

        if funcao_nome:
            funcao = funcoes.get(normalizar(funcao_nome))

            if not funcao:
                erros.append(f"Função '{funcao_nome}' não encontrada.")

        if cpf and Funcionario.objects.filter(cpf=cpf_original).exists():
            erros.append("CPF já cadastrado no sistema.")

        validade_cnh_original = valor_coluna(
            row,
            headers,
            "Validade CNH",
        )

        try:
            validade_cnh = normalizar_data(validade_cnh_original)
        except ValueError as exc:
            validade_cnh = None
            erros.append(str(exc))

        registros.append(
            {
                "linha": numero_linha,
                "nome": nome,
                "cpf": cpf,
                "email": texto(valor_coluna(row, headers, "E-mail")),
                "telefone": texto(valor_coluna(row, headers, "Telefone")),
                "numero_cnh": texto(valor_coluna(row, headers, "Número CNH")),
                "validade_cnh": validade_cnh,
                "funcao_id": funcao.pk if funcao else None,
                "funcao_nome": funcao.nome if funcao else "",
                "empresa_id": empresa.pk if empresa else None,
                "empresa_nome": (
                    empresa.nome_fantasia or empresa.razao_social
                    if empresa
                    else empresa_nome
                ),
                "cursos": texto(valor_coluna(row, headers, "Cursos")),
                "observacao": texto(valor_coluna(row, headers, "Observações")),
                "erros": erros,
            }
        )

    return registros


def ler_contas(worksheet):
    headers = obter_headers(worksheet)

    validar_colunas(
        headers,
        [
            "CPF Funcionário",
        ],
    )

    bancos = mapa_bancos()
    tipos_conta = mapa_tipo_conta()
    tipos_pix = mapa_tipo_pix()

    registros = []

    for numero_linha, row in enumerate(
        worksheet.iter_rows(min_row=2),
        start=2,
    ):
        cpf = normalizar_cpf(
            valor_coluna(
                row,
                headers,
                "CPF Funcionário",
            )
        )

        if not cpf:
            continue

        banco_original = texto(valor_coluna(row, headers, "Banco"))

        agencia = texto(valor_coluna(row, headers, "Agência"))

        digito_agencia = texto(valor_coluna(row, headers, "Dígito Agência"))

        numero_conta = texto(valor_coluna(row, headers, "Número Conta"))

        digito_conta = texto(valor_coluna(row, headers, "Dígito Conta"))

        tipo_conta_original = texto(valor_coluna(row, headers, "Tipo Conta"))

        tipo_pix_original = texto(valor_coluna(row, headers, "Tipo Chave PIX"))

        chave_pix = texto(valor_coluna(row, headers, "Chave PIX"))

        padrao = valor_bool(valor_coluna(row, headers, "Conta Padrão"))

        observacao = texto(valor_coluna(row, headers, "Observação"))

        # Verifica se a linha possui algum dado bancário.
        possui_dados = any(
            [
                banco_original,
                agencia,
                digito_agencia,
                numero_conta,
                digito_conta,
                tipo_conta_original,
                tipo_pix_original,
                chave_pix,
                padrao,
                observacao,
            ]
        )

        if not possui_dados:
            continue

        erros = []

        banco = bancos.get(normalizar(banco_original))

        if not banco:
            erros.append(f"Banco '{banco_original}' inválido.")

        if not agencia:
            erros.append("Agência não informada.")

        if not numero_conta:
            erros.append("Número da conta não informado.")

        tipo_conta = tipos_conta.get(normalizar(tipo_conta_original))

        if not tipo_conta:
            erros.append(f"Tipo de conta '{tipo_conta_original}' inválido.")

        tipo_pix = tipos_pix.get(normalizar(tipo_pix_original))

        if not tipo_pix:
            erros.append(f"Tipo de chave PIX '{tipo_pix_original}' inválido.")

        if not chave_pix:
            erros.append("Chave PIX não informada.")

        registros.append(
            {
                "linha": numero_linha,
                "cpf": cpf,
                "banco": banco,
                "agencia": agencia,
                "digito_agencia": digito_agencia,
                "numero_conta": numero_conta,
                "digito_conta": digito_conta,
                "tipo_conta": tipo_conta,
                "tipo_chave_pix": tipo_pix,
                "chave_pix": chave_pix,
                "padrao": padrao,
                "observacao": observacao,
                "erros": erros,
            }
        )

    return registros


def validar_relacionamento_contas(funcionarios, contas):
    cpfs_funcionarios = {
        normalizar_cpf(funcionario["cpf"])
        for funcionario in funcionarios
        if not funcionario["erros"]
    }

    contas_por_cpf = defaultdict(list)

    for conta in contas:
        contas_por_cpf[conta["cpf"]].append(conta)

        if conta["cpf"] not in cpfs_funcionarios:
            conta["erros"].append(
                "CPF da conta não corresponde a um funcionário "
                "válido da aba Funcionarios."
            )

    for cpf, lista in contas_por_cpf.items():
        padrao = [conta for conta in lista if conta["padrao"]]

        if len(padrao) > 1:
            for conta in padrao:
                conta["erros"].append(
                    "Existe mais de uma conta padrão para este funcionário."
                )


def analisar_arquivo(arquivo):
    try:
        workbook = load_workbook(
            filename=BytesIO(arquivo.read()),
            data_only=True,
        )
    except Exception as exc:
        raise ImportacaoFuncionarioError(
            "Não foi possível ler o arquivo Excel."
        ) from exc

    abas = set(workbook.sheetnames)

    if "Funcionarios" not in abas:
        raise ImportacaoFuncionarioError("A aba 'Funcionarios' não foi encontrada.")

    funcionarios = ler_funcionarios(workbook["Funcionarios"])

    contas = []

    if "ContasBancarias" in abas:
        contas = ler_contas(workbook["ContasBancarias"])

    validar_relacionamento_contas(
        funcionarios,
        contas,
    )

    total_funcionarios = len(funcionarios)
    funcionarios_com_erro = sum(bool(item["erros"]) for item in funcionarios)

    total_contas = len(contas)
    contas_com_erro = sum(bool(item["erros"]) for item in contas)

    return {
        "funcionarios": funcionarios,
        "contas": contas,
        "resumo": {
            "total_funcionarios": total_funcionarios,
            "funcionarios_validos": (total_funcionarios - funcionarios_com_erro),
            "funcionarios_com_erro": funcionarios_com_erro,
            "total_contas": total_contas,
            "contas_validas": (total_contas - contas_com_erro),
            "contas_com_erro": contas_com_erro,
        },
    }


@transaction.atomic
def executar_importacao(dados):
    funcionarios_criados = {}

    for item in dados["funcionarios"]:
        if item["erros"]:
            continue

        # Recupera as relações pelo ID.
        empresa = Empresa.objects.get(pk=item["empresa_id"])

        funcao = None

        if item["funcao_id"] is not None:
            funcao = Funcao.objects.get(pk=item["funcao_id"])

        validade_cnh = None

        if item["validade_cnh"]:
            validade_cnh = date.fromisoformat(item["validade_cnh"])

        funcionario = Funcionario.objects.create(
            nome=item["nome"],
            cpf=item["cpf"],
            email=item["email"],
            telefone=item["telefone"],
            numero_cnh=item["numero_cnh"],
            validade_cnh=validade_cnh,
            funcao=funcao,
            empresa=empresa,
            cursos=item["cursos"],
            observacao=item["observacao"],
        )

        funcionarios_criados[item["cpf"]] = funcionario

    contas_criadas = 0

    for item in dados["contas"]:
        if item["erros"]:
            continue

        funcionario = funcionarios_criados.get(item["cpf"])

        if funcionario is None:
            continue

        ContaBancariaFuncionario.objects.create(
            funcionario=funcionario,
            banco=item["banco"],
            agencia=item["agencia"],
            digito_agencia=item["digito_agencia"],
            numero_conta=item["numero_conta"],
            digito_conta=item["digito_conta"],
            tipo_conta=item["tipo_conta"],
            tipo_chave_pix=item["tipo_chave_pix"],
            chave_pix=item["chave_pix"],
            padrao=item["padrao"],
            observacao=item["observacao"],
        )

        contas_criadas += 1

    return {
        "funcionarios_criados": len(funcionarios_criados),
        "contas_criadas": contas_criadas,
    }
