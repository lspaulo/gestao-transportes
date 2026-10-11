from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from apps.cadastros.models import (
    ClasseOperacional,
    Empresa,
    StatusEquipamento,
)


def gerar_modelo_equipamentos():

    workbook = Workbook()

    ws = workbook.active

    if ws is None:
        raise RuntimeError("Não foi possível criar a planilha de equipamentos.")

    ws.title = "Equipamentos"

    colunas = [
        "Frota",
        "Empresa",
        "Classe Operacional",
        "Status Operacional",
        "Descrição",
        "Placa",
        "Marca",
        "Modelo",
        "Ano Fabricação",
        "Ano Modelo",
        "Renavam",
        "Chassi",
        "Cor",
        "Observação",
    ]

    ws.append(colunas)

    # Estilização do cabeçalho
    preenchimento = PatternFill(
        fill_type="solid",
        fgColor="198754",
    )

    for celula in ws[1]:
        celula.font = Font(
            bold=True,
            color="FFFFFF",
        )
        celula.fill = preenchimento
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:N1"

    larguras = {
        "A": 15,
        "B": 30,
        "C": 25,
        "D": 25,
        "E": 35,
        "F": 15,
        "G": 20,
        "H": 20,
        "I": 18,
        "J": 18,
        "K": 22,
        "L": 25,
        "M": 18,
        "N": 40,
    }

    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura

    # Aba auxiliar com opções válidas
    listas = workbook.create_sheet("Listas")

    listas.append(["Empresas", "Classes", "Status"])

    empresas = list(Empresa.objects.all().order_by("razao_social"))

    classes = list(ClasseOperacional.objects.all().order_by("nome"))

    status_equipamentos = list(StatusEquipamento.objects.all().order_by("nome"))

    for linha, empresa in enumerate(empresas, start=2):
        listas.cell(
            row=linha,
            column=1,
            value=empresa.nome_fantasia or empresa.razao_social,
        )

    for linha, classe in enumerate(classes, start=2):
        listas.cell(
            row=linha,
            column=2,
            value=classe.nome,
        )

    for linha, status in enumerate(status_equipamentos, start=2):
        listas.cell(
            row=linha,
            column=3,
            value=status.nome,
        )

    # Listas suspensas
    if empresas:
        validacao_empresas = DataValidation(
            type="list",
            formula1=f"'Listas'!$A$2:$A${len(empresas) + 1}",
            allow_blank=False,
        )
        ws.add_data_validation(validacao_empresas)
        validacao_empresas.add("B2:B1000")

    if classes:
        validacao_classes = DataValidation(
            type="list",
            formula1=f"'Listas'!$B$2:$B${len(classes) + 1}",
            allow_blank=False,
        )
        ws.add_data_validation(validacao_classes)
        validacao_classes.add("C2:C1000")

    if status_equipamentos:
        validacao_status = DataValidation(
            type="list",
            formula1=f"'Listas'!$C$2:$C${len(status_equipamentos) + 1}",
            allow_blank=False,
        )
        ws.add_data_validation(validacao_status)
        validacao_status.add("D2:D1000")

    # Formatos das colunas
    for linha in range(2, 1001):
        ws.cell(linha, 1).number_format = "@"
        ws.cell(linha, 6).number_format = "@"
        ws.cell(linha, 11).number_format = "@"
        ws.cell(linha, 12).number_format = "@"

    listas.sheet_state = "hidden"

    # Aba de instruções
    instrucoes = workbook.create_sheet("Instruções", 0)

    orientacoes = [
        ["IMPORTAÇÃO DE EQUIPAMENTOS"],
        [""],
        ["Campos obrigatórios:"],
        ["Frota"],
        ["Empresa"],
        ["Classe Operacional"],
        ["Status Operacional"],
        ["Descrição"],
        [""],
        ["Orientações:"],
        ["1. Não altere os nomes das colunas da aba Equipamentos."],
        ["2. Utilize as listas suspensas para empresa, classe e status."],
        ["3. A frota deve ser única e não pode estar cadastrada."],
        ["4. Preencha os anos de fabricação e modelo somente com números."],
        ["5. Deixe em branco os campos opcionais que não se aplicarem."],
        ["6. Não preencha as abas Listas e Instruções."],
        ["7. Revise a prévia de validação antes de confirmar a importação."],
    ]

    for linha in orientacoes:
        instrucoes.append(linha)

    instrucoes["A1"].font = Font(
        bold=True,
        size=14,
        color="198754",
    )

    instrucoes.column_dimensions["A"].width = 95

    for linha in range(1, instrucoes.max_row + 1):
        instrucoes.cell(linha, 1).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    return arquivo
